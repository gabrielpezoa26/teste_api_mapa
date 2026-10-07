"""Funcoes utilitarias: geometria, normalizacao, elevacao e escrita de planilhas."""
import math
import os
import re
import unicodedata

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import DATA, GAP_M, HYSTERESIS_M, SMOOTH_RADIUS


def fix_owner():
    try:
        st = os.stat(".")
        for f in [DATA, *DATA.rglob("*")]:
            os.chown(f, st.st_uid, st.st_gid)
    except (OSError, AttributeError):
        pass


# ---------- Geometria ----------
def haversine(a, b) -> float:
    (lat1, lon1), (lat2, lon2) = a, b
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = p2 - p1
    dl = math.radians(lon2 - lon1)
    h = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * 6371000 * math.asin(math.sqrt(h))


def line_length(coords) -> float:
    return sum(haversine(coords[i], coords[i + 1]) for i in range(len(coords) - 1))


def norm(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


class UF:
    def __init__(self, n):
        self.p = list(range(n))

    def find(self, x):
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        self.p[self.find(a)] = self.find(b)


def chain_ways(ways: list[dict]):
    remaining = sorted((list(w["coords"]) for w in ways), key=line_length, reverse=True)
    segments = []
    max_gap = 0.0

    while remaining:
        cur = remaining.pop(0)

        for end in ("tail", "head"):
            while True:
                point = cur[-1] if end == "tail" else cur[0]
                best = None

                for i, c in enumerate(remaining):
                    for first, inv in ((c[0], False), (c[-1], True)):
                        d = haversine(point, first)
                        if d <= GAP_M and (best is None or d < best[0]):
                            best = (d, i, inv)

                if best is None:
                    break

                d, i, inv = best
                piece = remaining.pop(i)

                if end == "tail":
                    cur = cur + (piece[::-1] if inv else piece)
                else:
                    cur = (piece if inv else piece[::-1]) + cur

                max_gap = max(max_gap, d)

        segments.append(cur)

    segments.sort(key=line_length, reverse=True)
    return segments, max_gap


# ---------- Elevacao ----------
def smooth(vals, r=SMOOTH_RADIUS):
    return [
        sum(vals[max(0, i - r):i + r + 1]) / len(vals[max(0, i - r):i + r + 1])
        for i in range(len(vals))
    ]


def gain_loss(vals, thr=HYSTERESIS_M):
    ref = vals[0]
    gain = 0.0
    loss = 0.0

    for e in vals[1:]:
        d = e - ref
        if d >= thr:
            gain += d
            ref = e
        elif d <= -thr:
            loss += -d
            ref = e

    return gain, loss


# ---------- Excel ----------
def write_sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)

    for c in ws[1]:
        c.font = Font(bold=True)

    for r in rows:
        ws.append(r)

    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)

    ws.freeze_panes = "A2"