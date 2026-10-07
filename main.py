"""Overpass -> junta ways por nome -> ordena -> elevacao -> dificuldade -> Excel + GeoJSON + GPX + SQLite.

Objetivo: base BRUTA. Nada é descartado na coleta (todos os ways nomeados de path/footway/track).
"""
import atexit
import json
import sqlite3
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook

from config import (
    BACKOFF_S,
    BBOX,
    CIRCULAR_M,
    CLIMB_SPEED_M_H,
    DEM_FILE,
    EASY_MAX,
    EFFORT_DESNIVEL_DIV,
    ELEVATION,
    ENDPOINTS,
    EXCEL_HEADERS,
    EXTRA_TAGS,
    GAP_M,
    GPX_DIR,
    HEADERS,
    HIGHWAYS,
    MAX_ATTEMPTS,
    MODERATE_MAX,
    NAME_RE,
    OUT_EXCEL,
    OUT_GEOJSON,
    OUT_SQLITE,
    OVERPASS_QUERY_TIMEOUT_S,
    QUERY_PAUSE_S,
    RAW_GLOB,
    REFRESH,
    REQUEST_TIMEOUT_S,
    RETRY_SLEEP_S,
    RUN_DATE,
    SAC_ORDER,
    SQLITE_BASE_COLUMNS,
    UTILS_DIR,
    WALK_SPEED_KMH,
)
from utils import (
    UF,
    chain_ways,
    fix_owner,
    gain_loss,
    haversine,
    line_length,
    norm,
    smooth,
    write_sheet,
)

atexit.register(fix_owner)


# ---------- 1. Overpass ----------
def build_query(highway: str) -> str:
    return (
        f"[out:json][timeout:{OVERPASS_QUERY_TIMEOUT_S}];\n"
        f"way[\"highway\"=\"{highway}\"][name]({BBOX});\n"
        f"out meta geom;"
    )


def post_overpass(query: str):
    last_err = None

    for attempt in range(1, MAX_ATTEMPTS + 1):
        for url in ENDPOINTS:
            try:
                print(f"  [{attempt}/{MAX_ATTEMPTS}] {url} ...")
                r = requests.post(
                    url,
                    data={"data": query},
                    headers=HEADERS,
                    timeout=REQUEST_TIMEOUT_S,
                )
                if r.status_code == 200:
                    return r.json(), url
                last_err = f"HTTP {r.status_code}"
            except Exception as e:
                last_err = str(e)

            print(f"    falhou: {last_err}")
            time.sleep(RETRY_SLEEP_S)

        time.sleep(BACKOFF_S * attempt)

    sys.exit(f"Overpass indisponivel: {last_err}.")


def fetch() -> tuple[dict, dict]:
    UTILS_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(UTILS_DIR.glob(RAW_GLOB))

    if files and not REFRESH:
        saved = json.loads(files[-1].read_text(encoding="utf-8"))
        meta = saved.get("meta", {})
        if meta.get("bbox") == BBOX and meta.get("highways") == HIGHWAYS:
            meta["cache_file"] = files[-1].name
            print(f"Usando cache {files[-1]}")
            return saved, meta

    elements = {}
    endpoints = {}
    osm_bases = []

    for i, hw in enumerate(HIGHWAYS):
        if i:
            time.sleep(QUERY_PAUSE_S)

        query = build_query(hw)
        print(f"Query highway={hw}...")
        data, url = post_overpass(query)
        endpoints[hw] = url

        if base := data.get("osm3s", {}).get("timestamp_osm_base", ""):
            osm_bases.append(base)

        for el in data.get("elements", []):
            elements[(el.get("type"), el.get("id"))] = el

    meta = {
        "extracted_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "osm_base": min(osm_bases) if osm_bases else "",
        "bbox": BBOX,
        "highways": HIGHWAYS,
        "endpoints": endpoints,
    }
    saved = {"meta": meta, "elements": list(elements.values())}

    out = UTILS_DIR / f"overpass_raw_{RUN_DATE}.json"
    out.write_text(json.dumps(saved), encoding="utf-8")
    meta["cache_file"] = out.name
    return saved, meta


# ---------- 2. Ways e agrupamento ----------
def parse_ways(data: dict) -> list[dict]:
    ways = []

    for el in data.get("elements", []):
        if el.get("type") != "way" or not el.get("geometry"):
            continue

        tags = el.get("tags", {})
        coords = [(p["lat"], p["lon"]) for p in el["geometry"]]
        ways.append({
            "id": el["id"],
            "name": tags.get("name", ""),
            "key": norm(tags.get("name", "")),
            "highway": tags.get("highway", ""),
            "tags": tags,
            "timestamp": el.get("timestamp", ""),
            "version": el.get("version", 0),
            "nodes": el.get("nodes", []),
            "coords": coords,
            "length_m": line_length(coords),
        })

    return ways


def merge_group(group: list[dict]) -> list[dict]:
    n = len(group)
    uf_nodes = UF(n)
    owner = {}

    for i, w in enumerate(group):
        for nid in w["nodes"]:
            if nid in owner:
                uf_nodes.union(i, owner[nid])
            else:
                owner[nid] = i

    node_root = [uf_nodes.find(i) for i in range(n)]

    uf_gap = UF(n)
    for i in range(n):
        uf_gap.union(i, node_root[i])

    ends = []
    for i, w in enumerate(group):
        ends.extend([(w["coords"][0], i), (w["coords"][-1], i)])

    for a in range(len(ends)):
        for b in range(a + 1, len(ends)):
            ia = ends[a][1]
            ib = ends[b][1]
            if uf_gap.find(ia) != uf_gap.find(ib) and haversine(ends[a][0], ends[b][0]) <= GAP_M:
                uf_gap.union(ia, ib)

    clusters = defaultdict(list)
    for i in range(n):
        clusters[uf_gap.find(i)].append(i)

    result = []
    for idxs in clusters.values():
        ws = [group[i] for i in idxs]
        pts = [p for w in ws for p in w["coords"]]
        stamps = [w["timestamp"] for w in ws if w["timestamp"]]
        segments, max_gap = chain_ways(ws)
        main = segments[0]

        all_tags = {}
        for w in ws:
            for k, v in w["tags"].items():
                all_tags.setdefault(k, set()).add(v)

        way_ids_list = [w["id"] for w in ws]
        trilha_id = f"TRL-{min(way_ids_list)}"  # ID unico e estavel baseado no way original

        t = {
            "trilha_id": trilha_id,
            "name": ws[0]["name"],
            "name_matches": "sim" if (NAME_RE is None or NAME_RE.search(ws[0]["name"])) else "nao",
            "n_ways": len(ws),
            "pieces_by_node": len({node_root[i] for i in idxs}),
            "segments": segments,
            "n_segments": len(segments),
            "max_gap_m": round(max_gap),
            "length_km": round(sum(w["length_m"] for w in ws) / 1000, 2),
            "km_principal": round(line_length(main) / 1000, 2),
            "lat": round(sum(p[0] for p in pts) / len(pts), 5),
            "lon": round(sum(p[1] for p in pts) / len(pts), 5),
            "start": main[0],
            "end": main[-1],
            "circular": haversine(main[0], main[-1]) < CIRCULAR_M,
            "highway": ", ".join(sorted({w["highway"] for w in ws if w["highway"]})),
            "last_edit": max(stamps) if stamps else "",
            "first_edit": min(stamps) if stamps else "",
            "max_version": max(w["version"] for w in ws),
            "tags_json": json.dumps(
                {k: sorted(v) for k, v in sorted(all_tags.items())},
                ensure_ascii=False,
            )[:32000],
            "way_ids_list": way_ids_list,
            "elev": None,
            "elev_segments": [],
            "elev_status": "desativada",
        }

        for k in EXTRA_TAGS:
            t[k] = ", ".join(sorted({w["tags"][k] for w in ws if w["tags"].get(k)}))

        result.append(t)

    return result


# ---------- 3. Elevacao OFFLINE (Copernicus DEM) ----------
def add_elevation_offline(trails):
    try:
        import rasterio
    except ImportError:
        print("Aviso: 'rasterio' não instalado (pip install rasterio). Elevação ignorada.")
        for t in trails:
            t["elev_status"] = "falhou (sem rasterio)"
        return

    if not Path(DEM_FILE).exists():
        print(f"Aviso: DEM {DEM_FILE} não encontrado. Elevação ignorada.")
        for t in trails:
            t["elev_status"] = "falhou (sem DEM)"
        return

    print(f"Calculando elevação offline em cada nó via {DEM_FILE}...")
    with rasterio.open(DEM_FILE) as src:
        for t in trails:
            all_pts = []
            for seg in t["segments"]:
                # Rasterio sample exige (lon, lat)
                all_pts.extend([(p[1], p[0]) for p in seg])

            try:
                # Extrai a elevacao instantaneamente para todos os pontos da trilha
                elevs = [float(val[0]) for val in src.sample(all_pts)]
            except Exception:
                elevs = [0.0] * len(all_pts)

            seg_elevs = []
            idx = 0
            for seg in t["segments"]:
                seg_elevs.append(elevs[idx:idx + len(seg)])
                idx += len(seg)
            t["elev_segments"] = seg_elevs

            # Ganho, perda e desnível baseados EXCLUSIVAMENTE no km_principal
            main_elev = smooth(seg_elevs[0]) if seg_elevs[0] else []
            gain, loss = gain_loss(main_elev) if main_elev else (0.0, 0.0)
            desnivel_acumulado = max(gain, loss)

            t["elev"] = {
                "gain": round(gain),
                "loss": round(loss),
                "desnivel_acumulado": round(desnivel_acumulado),
                "min": round(min(elevs)) if elevs else 0,
                "max": round(max(elevs)) if elevs else 0,
            }
            t["elev_status"] = "ok"


# ---------- 4. Dificuldade e Tempo ----------
def classify(t: dict):
    sac = [s.strip() for s in t["sac_scale"].split(",") if s.strip() in SAC_ORDER]
    if sac:
        worst = max(SAC_ORDER.index(s) for s in sac)
        level = "facil" if worst == 0 else "moderada" if worst == 1 else "dificil"
        return level, "osm"

    # Dificuldade usa o desnível acumulado como proxy de esforço
    desnivel = t["elev"]["desnivel_acumulado"] if t["elev"] else 0
    effort = t["km_principal"] + (desnivel / EFFORT_DESNIVEL_DIV)
    level = "facil" if effort <= EASY_MAX else "moderada" if effort <= MODERATE_MAX else "dificil"
    return level, "calculado" if t["elev"] else "calculado_sem_desnivel"


def est_time(t: dict):
    desnivel = t["elev"]["desnivel_acumulado"] if t["elev"] else 0
    minutes = round((t["km_principal"] / WALK_SPEED_KMH + desnivel / CLIMB_SPEED_M_H) * 60)
    return minutes, f"{minutes // 60}h{minutes % 60:02d}"


# ---------- 5. Saídas (GeoJSON, GPX e SQLite) ----------
def write_geojson_file(trails):
    OUT_GEOJSON.parent.mkdir(parents=True, exist_ok=True)
    feats = []

    for t in trails:
        el = t["elev"] or {}
        feats.append({
            "type": "Feature",
            "geometry": {
                "type": "MultiLineString",
                "coordinates": [
                    [[round(p[1], 6), round(p[0], 6)] for p in s]
                    for s in t["segments"]
                ],
            },
            "properties": {
                "trilha_id": t["trilha_id"],
                "name": t["name"],
                "km_total": t["length_km"],
                "km_principal": t["km_principal"],
                "elevation_status": t["elev_status"],
                "difficulty_calc": t["difficulty"],
                "desnivel_m": el.get("desnivel_acumulado"),
                "gain_m": el.get("gain"),
                "loss_m": el.get("loss"),
                "time_calc": t["time"],
            },
        })

    OUT_GEOJSON.write_text(
        json.dumps({"type": "FeatureCollection", "features": feats}, ensure_ascii=False),
        encoding="utf-8",
    )


def write_gpx_file(trails):
    GPX_DIR.mkdir(parents=True, exist_ok=True)

    for t in trails:
        lines = [
            '<?xml version="1.0" encoding="UTF-8"?>',
            '<gpx version="1.1" creator="Trilhas SP" xmlns="http://www.topografix.com/GPX/1/1">',
            f'  <trk>\n    <name>{t["name"]}</name>',
        ]

        for i, seg in enumerate(t["segments"]):
            lines.append('    <trkseg>')
            seg_elevs = t["elev_segments"][i] if t["elev_segments"] else None

            for j, p in enumerate(seg):
                ele_str = f"<ele>{seg_elevs[j]:.1f}</ele>" if seg_elevs else ""
                lines.append(f'      <trkpt lat="{p[0]:.6f}" lon="{p[1]:.6f}">{ele_str}</trkpt>')

            lines.append('    </trkseg>')

        lines.append('  </trk>\n</gpx>')

        # Ex: TRL-12345678_20241026.gpx
        out_file = GPX_DIR / f"{t['trilha_id']}_{RUN_DATE}.gpx"
        out_file.write_text("\n".join(lines), encoding="utf-8")


def write_sqlite_file(trails, ways):
    OUT_SQLITE.parent.mkdir(parents=True, exist_ok=True)

    if OUT_SQLITE.exists():
        OUT_SQLITE.unlink()

    conn = sqlite3.connect(OUT_SQLITE)
    c = conn.cursor()

    # Adiciona dinamicamente as extra_tags
    extra_columns = [f"{tag} TEXT" for tag in EXTRA_TAGS]

    # Junta tudo e mantém a geometria GeoJSON no final
    all_columns = SQLITE_BASE_COLUMNS + extra_columns + ["geom_geojson TEXT"]

    # 2. Criação das tabelas
    create_table_sql = f"CREATE TABLE trilhas (\n    {',\n    '.join(all_columns)}\n)"
    c.execute(create_table_sql)

    c.execute('''CREATE TABLE trilha_ways (
        trilha_id TEXT, way_id INTEGER,
        FOREIGN KEY(trilha_id) REFERENCES trilhas(trilha_id)
    )''')

    c.execute('''CREATE TABLE ways (
        way_id INTEGER PRIMARY KEY, nome TEXT, highway TEXT, ultima_edicao TEXT
    )''')

    for w in ways:
        c.execute(
            "INSERT OR IGNORE INTO ways VALUES (?, ?, ?, ?)",
            (w["id"], w["name"], w["highway"], w["timestamp"]),
        )

    # 3. Inserção dos dados
    num_cols = len(all_columns)
    placeholders = ", ".join(["?"] * num_cols)
    insert_sql = f"INSERT INTO trilhas VALUES ({placeholders})"

    for t in trails:
        el = t["elev"] or {}
        geom = json.dumps({
            "type": "MultiLineString",
            "coordinates": [[[p[1], p[0]] for p in s] for s in t["segments"]],
        })

        # Monta a tupla de valores exatamente como na aba de trilhas do Excel
        row = [
            t["trilha_id"],
            t["name"],
            t["name_matches"],
            t["km_principal"],
            t["length_km"],
            t["n_ways"],
            t["n_segments"],
            round(t["start"][0], 5),
            round(t["start"][1], 5),
            round(t["end"][0], 5),
            round(t["end"][1], 5),
            "sim" if t["circular"] else "nao",  # Formato texto como no Excel
            t["difficulty"],
            t["difficulty_source"],
            t["elev_status"],
            el.get("desnivel_acumulado"),
            el.get("gain"),
            el.get("loss"),
            t["time"],
            t["last_edit"],
            t["highway"],
        ]

        # Adiciona os valores das extra tags
        row.extend([t[k] for k in EXTRA_TAGS])

        # Adiciona a geometria por último
        row.append(geom)

        c.execute(insert_sql, row)

        # Popula a tabela de relacionamento
        for wid in t["way_ids_list"]:
            c.execute("INSERT INTO trilha_ways VALUES (?, ?)", (t["trilha_id"], wid))

    conn.commit()
    conn.close()


def write_excel_file(trails, ways, meta):
    OUT_EXCEL.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    write_sheet(
        wb,
        "resumo",
        ["metrica", "valor"],
        [
            ["data da extracao (UTC)", meta.get("extracted_at", "")],
            ["arquivo sqlite gerado", OUT_SQLITE.name],
            ["diretorio gpx", GPX_DIR.name],
            ["total de trilhas geradas", len(trails)],
        ],
        first=True,
    )

    rows = []
    for t in trails:
        el = t["elev"] or {}
        rows.append([
            t["trilha_id"], t["name"], t["name_matches"], t["km_principal"],
            t["length_km"], t["n_ways"], t["n_segments"],
            round(t["start"][0], 5), round(t["start"][1], 5),
            round(t["end"][0], 5), round(t["end"][1], 5),
            "sim" if t["circular"] else "nao", t["difficulty"], t["difficulty_source"],
            t["elev_status"], el.get("desnivel_acumulado"), el.get("gain"), el.get("loss"),
            t["time"], t["last_edit"], t["highway"],
            *[t[k] for k in EXTRA_TAGS],
        ])
    write_sheet(wb, "trilhas", EXCEL_HEADERS, rows)

    # Aba relacional trilha_ways
    t_ways_rows = []
    for t in trails:
        for wid in t["way_ids_list"]:
            t_ways_rows.append([t["trilha_id"], wid])
    write_sheet(wb, "trilha_ways", ["trilha_id", "way_id"], t_ways_rows)

    # Aba ways brutos
    write_sheet(
        wb,
        "ways_brutos",
        ["way_id", "nome", "highway", "metros", "ultima_edicao"],
        [[w["id"], w["name"], w["highway"], round(w["length_m"]), w["timestamp"]] for w in ways],
    )

    wb.save(OUT_EXCEL)


def main():
    data, meta = fetch()
    ways = parse_ways(data)
    print(f"{len(ways)} ways recebidos")

    groups = defaultdict(list)
    for w in ways:
        if w["key"]:
            groups[w["key"]].append(w)

    trails = []
    for g in groups.values():
        trails.extend(merge_group(g))
    trails.sort(key=lambda t: t["km_principal"], reverse=True)

    if ELEVATION:
        add_elevation_offline(trails)

    for t in trails:
        t["difficulty"], t["difficulty_source"] = classify(t)
        t["time_min"], t["time"] = est_time(t)

    write_excel_file(trails, ways, meta)
    write_geojson_file(trails)
    write_gpx_file(trails)
    write_sqlite_file(trails, ways)

    print(f"\nOK -> {OUT_EXCEL}, {OUT_GEOJSON}, e {OUT_SQLITE}")
    print(f"GPXs salvos em -> {GPX_DIR}/")


if __name__ == "__main__":
    main()